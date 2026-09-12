"""
Reports & Export Service
========================
Generates professional PDF and Excel reports for all roles in the skincare platform.
Uses reportlab for PDF and openpyxl for Excel generation.
"""

import io
from datetime import datetime, timezone, date
from typing import Optional, List, Dict, Any

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, mm
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    PageBreak, HRFlowable
)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ═══════════════════════════════════════════════════════════════════════════════
# PDF Helpers
# ═══════════════════════════════════════════════════════════════════════════════

BRAND_COLOR = colors.HexColor("#6C3CE1")
BRAND_DARK = colors.HexColor("#1E1B3A")
ACCENT_COLOR = colors.HexColor("#14B8A6")
RISK_COLORS = {
    "Critical": colors.HexColor("#DC2626"),
    "High": colors.HexColor("#EA580C"),
    "Medium": colors.HexColor("#D97706"),
    "Low": colors.HexColor("#16A34A"),
}
CATEGORY_COLORS = {
    "Excellent": colors.HexColor("#16A34A"),
    "Good": colors.HexColor("#22C55E"),
    "Fair": colors.HexColor("#EAB308"),
    "Poor": colors.HexColor("#DC2626"),
}


def _get_styles():
    """Build custom paragraph styles for PDF reports."""
    styles = getSampleStyleSheet()

    styles.add(ParagraphStyle(
        "BrandTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        textColor=BRAND_DARK,
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        "BrandSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=11,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=14,
    ))
    styles.add(ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=14,
        textColor=BRAND_COLOR,
        spaceBefore=18,
        spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        "SubSection",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=11,
        textColor=BRAND_DARK,
        spaceBefore=12,
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        "CellText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
    ))
    styles.add(ParagraphStyle(
        "FooterStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        textColor=colors.HexColor("#94A3B8"),
        alignment=TA_CENTER,
    ))
    return styles


def _build_header(styles, title: str, subtitle: str, role_label: str = ""):
    """Build the report header elements."""
    elements = []
    elements.append(Paragraph("AI Skin Intelligence", styles["BrandSubtitle"]))
    elements.append(Paragraph(title, styles["BrandTitle"]))
    meta_parts = [subtitle]
    if role_label:
        meta_parts.append(f"Role: {role_label}")
    meta_parts.append(f"Generated: {datetime.now(timezone.utc).strftime('%B %d, %Y at %H:%M UTC')}")
    elements.append(Paragraph(" • ".join(meta_parts), styles["BrandSubtitle"]))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=BRAND_COLOR, spaceAfter=12))
    return elements


def _build_kv_table(data: List[tuple], col_widths=None):
    """Build a two-column key-value table."""
    if not data:
        return Spacer(1, 0)

    if col_widths is None:
        col_widths = [2.2 * inch, 4.3 * inch]

    table = Table(data, colWidths=col_widths)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F8F5FF")),
        ("TEXTCOLOR", (0, 0), (0, -1), BRAND_DARK),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
    ]))
    return table


def _build_data_table(headers: List[str], rows: List[list], col_widths=None):
    """Build a multi-column data table with header styling."""
    if not rows:
        return Spacer(1, 0)

    table_data = [headers] + rows
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), BRAND_COLOR),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]
    table.setStyle(TableStyle(style_commands))
    return table


def _build_footer():
    """Build report footer."""
    styles = _get_styles()
    return Paragraph(
        f"AI Skin Intelligence Platform — Confidential Report — {date.today().isoformat()}",
        styles["FooterStyle"]
    )


def _safe_str(val, default="—"):
    """Safely convert value to string for display."""
    if val is None or (isinstance(val, str) and not val.strip()):
        return default
    return str(val)


# ═══════════════════════════════════════════════════════════════════════════════
# Excel Helpers
# ═══════════════════════════════════════════════════════════════════════════════

EXCEL_HEADER_FILL = PatternFill(start_color="6C3CE1", end_color="6C3CE1", fill_type="solid")
EXCEL_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
EXCEL_TITLE_FONT = Font(name="Calibri", bold=True, color="1E1B3A", size=14)
EXCEL_SUBTITLE_FONT = Font(name="Calibri", color="64748B", size=10)
EXCEL_CELL_FONT = Font(name="Calibri", size=10)
EXCEL_THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
RISK_FILLS = {
    "Critical": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "High": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
    "Medium": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    "Low": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
}


def _create_workbook_with_title(title: str, subtitle: str = ""):
    """Create an openpyxl workbook with a branded title header."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.merge_cells("A1:F1")
    ws["A1"] = f"AI Skin Intelligence — {title}"
    ws["A1"].font = EXCEL_TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left")

    ws.merge_cells("A2:F2")
    ws["A2"] = subtitle or f"Generated: {datetime.now(timezone.utc).strftime('%B %d, %Y %H:%M UTC')}"
    ws["A2"].font = EXCEL_SUBTITLE_FONT
    return wb, ws


def _write_headers(ws, row: int, headers: List[str]):
    """Write styled header row."""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = EXCEL_HEADER_FONT
        cell.fill = EXCEL_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = EXCEL_THIN_BORDER


def _write_data_rows(ws, start_row: int, data: List[List[Any]], risk_col: int = None):
    """Write data rows with optional risk-level coloring."""
    for row_idx, row_data in enumerate(data, start_row):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = EXCEL_CELL_FONT
            cell.border = EXCEL_THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if risk_col and col_idx == risk_col:
                fill = RISK_FILLS.get(str(value), None)
                if fill:
                    cell.fill = fill


def _auto_fit_columns(ws, min_width=12, max_width=50):
    """Auto-fit column widths based on content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                cell_len = len(str(cell.value))
                if cell_len > max_len:
                    max_len = cell_len
        adjusted = max(min_width, min(max_len + 4, max_width))
        ws.column_dimensions[col_letter].width = adjusted


# ═══════════════════════════════════════════════════════════════════════════════
# Data Builders — Gather data from DB models
# ═══════════════════════════════════════════════════════════════════════════════

def _get_val(obj, key, default=""):
    """Safely get an attribute or dict key from an object."""
    if obj is None:
        return default
    if isinstance(obj, dict):
        return obj.get(key, default)
    return getattr(obj, key, default)


def _get_user_profile_data(user, profile):
    """Build common profile key-value data."""
    u_name = _get_val(user, "name", "")
    if not u_name:
        u_name = f"{_get_val(user, 'first_name', '')} {_get_val(user, 'last_name', '')}".strip()
    return [
        ("Name", _safe_str(u_name or _get_val(user, "email", ""))),
        ("Email", _safe_str(_get_val(user, "email", ""))),
        ("Skin Type", _safe_str(_get_val(profile, "skin_type", None))),
        ("Age Group", _safe_str(_get_val(profile, "age_group", None))),
        ("Skin Concerns", _safe_str(_get_val(profile, "skin_concerns", None))),
        ("Allergies", _safe_str(_get_val(profile, "allergies", None))),
        ("Sensitivities", _safe_str(_get_val(profile, "sensitivities", None))),
        ("Lifestyle Habits", _safe_str(_get_val(profile, "lifestyle_habits", None))),
        ("Sleep Quality", _safe_str(_get_val(profile, "sleep_quality", None))),
        ("Water Intake", _safe_str(_get_val(profile, "water_intake", None))),
        ("Environmental Exposure", _safe_str(_get_val(profile, "environmental_exposure", None))),
        ("Skin Health Score", str(_get_val(profile, "skin_health_score", 0))),
    ]


def _fmt_dt(dt_val, fmt="%Y-%m-%d %H:%M"):
    """Format a datetime safely."""
    if dt_val:
        try:
            return dt_val.strftime(fmt)
        except Exception:
            return str(dt_val)
    return "—"


# ═══════════════════════════════════════════════════════════════════════════════
# REPORT GENERATORS — Each returns (pdf_bytes | excel_bytes)
# ═══════════════════════════════════════════════════════════════════════════════


def build_skin_assessment_report(user, profile, assessments, format: str = "pdf") -> bytes:
    """
    Skin Assessment Report: profile info, latest score, risks, priority concerns.
    """
    latest = assessments[0] if assessments else None
    user_name = _get_val(user, "name") or _get_val(user, "email")

    if format == "excel":
        wb, ws = _create_workbook_with_title(
            "Skin Assessment Report",
            f"Patient: {user_name} | Generated: {datetime.now(timezone.utc).strftime('%B %d, %Y %H:%M UTC')}"
        )

        # Profile sheet
        ws.title = "Skin Profile"
        row = 4
        profile_data = _get_user_profile_data(user, profile)
        _write_headers(ws, row, ["Field", "Value"])
        row += 1
        _write_data_rows(ws, row, profile_data)
        _auto_fit_columns(ws)

        # Assessment sheet
        ws2 = wb.create_sheet("Assessment History")
        ws2.merge_cells("A1:G1")
        ws2["A1"] = "Assessment History"
        ws2["A1"].font = EXCEL_TITLE_FONT
        _write_headers(ws2, 3, ["#", "Date", "Score", "Category", "Risk Level", "Trigger", "Notes"])
        rows_data = []
        for idx, a in enumerate(assessments[:50], 1):
            rows_data.append([
                idx, _fmt_dt(_get_val(a, "assessment_date")), _get_val(a, "skin_health_score", 0),
                _get_val(a, "skin_health_category", "—"), _get_val(a, "overall_risk_level", "—"),
                _get_val(a, "trigger_source", "—"), _safe_str(_get_val(a, "notes", ""))
            ])
        _write_data_rows(ws2, 4, rows_data, risk_col=5)
        _auto_fit_columns(ws2)

        # Risks sheet
        risks = _get_val(latest, "risks", []) if latest else []
        if risks:
            ws3 = wb.create_sheet("Risk Factors")
            ws3.merge_cells("A1:D1")
            ws3["A1"] = "Identified Risk Factors"
            ws3["A1"].font = EXCEL_TITLE_FONT
            _write_headers(ws3, 3, ["Risk Title", "Risk Level", "Description", "Recommendation"])
            risk_rows = [[_get_val(r, "risk_title", ""), _get_val(r, "risk_level", ""), _get_val(r, "description", ""), _get_val(r, "recommendation", "")] for r in risks]
            _write_data_rows(ws3, 4, risk_rows, risk_col=2)
            _auto_fit_columns(ws3)

        # Priorities sheet
        priorities = _get_val(latest, "priorities", []) if latest else []
        if priorities:
            ws4 = wb.create_sheet("Priority Concerns")
            ws4.merge_cells("A1:D1")
            ws4["A1"] = "Prioritized Skin Concerns"
            ws4["A1"].font = EXCEL_TITLE_FONT
            _write_headers(ws4, 3, ["Concern", "Rank", "Severity", "Score"])
            prio_rows = [[_get_val(p, "concern_name", ""), _get_val(p, "priority_rank", ""), _get_val(p, "severity", ""), _get_val(p, "priority_score", "")] for p in priorities]
            _write_data_rows(ws4, 4, prio_rows, risk_col=3)
            _auto_fit_columns(ws4)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "Skin Assessment Report", f"Patient: {user_name}")

    # Profile section
    elements.append(Paragraph("Skin Profile", styles["SectionHeading"]))
    profile_kv = _get_user_profile_data(user, profile)
    elements.append(_build_kv_table(profile_kv))

    # Latest assessment
    if latest:
        elements.append(Spacer(1, 12))
        elements.append(Paragraph("Latest Assessment", styles["SectionHeading"]))
        assessment_kv = [
            ("Health Score", str(_get_val(latest, "skin_health_score", 0))),
            ("Category", str(_get_val(latest, "skin_health_category", "Fair"))),
            ("Overall Risk Level", str(_get_val(latest, "overall_risk_level", "Low"))),
            ("Assessment Date", _fmt_dt(_get_val(latest, "assessment_date", None))),
            ("Trigger Source", str(_get_val(latest, "trigger_source", ""))),
            ("Notes", _safe_str(_get_val(latest, "notes", ""))),
        ]
        elements.append(_build_kv_table(assessment_kv))

        # Risks
        risks = _get_val(latest, "risks", [])
        if risks:
            elements.append(Spacer(1, 10))
            elements.append(Paragraph("Identified Risk Factors", styles["SubSection"]))
            risk_headers = ["Risk Title", "Level", "Description", "Recommendation"]
            risk_rows = [[_get_val(r, "risk_title", ""), _get_val(r, "risk_level", ""), str(_get_val(r, "description", ""))[:80], str(_get_val(r, "recommendation", ""))[:80]] for r in risks]
            elements.append(_build_data_table(risk_headers, risk_rows, col_widths=[1.3*inch, 0.7*inch, 2.3*inch, 2.3*inch]))

        # Priorities
        priorities = _get_val(latest, "priorities", [])
        if priorities:
            elements.append(Spacer(1, 10))
            elements.append(Paragraph("Priority Concerns", styles["SubSection"]))
            prio_headers = ["Concern", "Rank", "Severity", "Score"]
            prio_rows = [[_get_val(p, "concern_name", ""), str(_get_val(p, "priority_rank", "")), _get_val(p, "severity", ""), str(_get_val(p, "priority_score", ""))] for p in priorities]
            elements.append(_build_data_table(prio_headers, prio_rows, col_widths=[2.5*inch, 1*inch, 1.5*inch, 1.5*inch]))

    # Assessment history table
    if len(assessments) > 1:
        elements.append(Spacer(1, 12))
        elements.append(Paragraph("Assessment History", styles["SectionHeading"]))
        hist_headers = ["#", "Date", "Score", "Category", "Risk Level"]
        hist_rows = [
            [str(i+1), _fmt_dt(_get_val(a, "assessment_date")), str(_get_val(a, "skin_health_score", 0)), str(_get_val(a, "skin_health_category", "")), str(_get_val(a, "overall_risk_level", ""))]
            for i, a in enumerate(assessments[:20])
        ]
        elements.append(_build_data_table(hist_headers, hist_rows, col_widths=[0.5*inch, 1.8*inch, 1*inch, 1.5*inch, 1.5*inch]))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_routine_report(user, routine, format: str = "pdf") -> bytes:
    """
    Routine Report: morning/evening/weekly steps, seasonal recommendations.
    """
    user_name = _get_val(user, "name") or _get_val(user, "email")
    steps = _get_val(routine, "steps", [])
    if not steps:
        m_list = _get_val(routine, "morning_steps", []) or []
        e_list = _get_val(routine, "evening_steps", []) or []
        w_list = _get_val(routine, "weekly_steps", []) or []
        combined = []
        for s in m_list:
            item = dict(s) if isinstance(s, dict) else s
            if isinstance(item, dict):
                item["time_of_day"] = "morning"
            combined.append(item)
        for s in e_list:
            item = dict(s) if isinstance(s, dict) else s
            if isinstance(item, dict):
                item["time_of_day"] = "evening"
            combined.append(item)
        for s in w_list:
            item = dict(s) if isinstance(s, dict) else s
            if isinstance(item, dict):
                item["time_of_day"] = "weekly"
            combined.append(item)
        steps = combined

    morning = [s for s in steps if _get_val(s, "time_of_day") == "morning" and _get_val(s, "is_active", True)]
    evening = [s for s in steps if _get_val(s, "time_of_day") == "evening" and _get_val(s, "is_active", True)]
    weekly = [s for s in steps if _get_val(s, "time_of_day") == "weekly" and _get_val(s, "is_active", True)]
    seasonal = _get_val(routine, "seasonal_recommendations", []) if routine else []

    if format == "excel":
        wb, ws = _create_workbook_with_title(
            "Skincare Routine Report",
            f"Patient: {user_name} | Season: {_get_val(routine, 'season', '—')}"
        )
        ws.title = "Morning Routine"
        _write_headers(ws, 4, ["#", "Step", "Category", "Description", "Active Ingredients", "Frequency", "Caution Notes"])
        rows = [[i+1, _get_val(s, "step_title", _get_val(s, "product_name", "")), _get_val(s, "category", _get_val(s, "product_type", "")), str(_get_val(s, "description", _get_val(s, "instructions", "")))[:100], _get_val(s, "active_ingredients", _get_val(s, "key_actives", "")) or "—", _get_val(s, "frequency", "Daily"), _get_val(s, "caution_notes", "") or "—"] for i, s in enumerate(morning)]
        _write_data_rows(ws, 5, rows)
        _auto_fit_columns(ws)

        ws2 = wb.create_sheet("Evening Routine")
        _write_headers(ws2, 1, ["#", "Step", "Category", "Description", "Active Ingredients", "Frequency", "Caution Notes"])
        rows2 = [[i+1, _get_val(s, "step_title", _get_val(s, "product_name", "")), _get_val(s, "category", _get_val(s, "product_type", "")), str(_get_val(s, "description", _get_val(s, "instructions", "")))[:100], _get_val(s, "active_ingredients", _get_val(s, "key_actives", "")) or "—", _get_val(s, "frequency", "Daily"), _get_val(s, "caution_notes", "") or "—"] for i, s in enumerate(evening)]
        _write_data_rows(ws2, 2, rows2)
        _auto_fit_columns(ws2)

        ws3 = wb.create_sheet("Weekly Treatments")
        _write_headers(ws3, 1, ["#", "Step", "Category", "Description", "Active Ingredients", "Frequency", "Caution Notes"])
        rows3 = [[i+1, _get_val(s, "step_title", _get_val(s, "product_name", "")), _get_val(s, "category", _get_val(s, "product_type", "")), str(_get_val(s, "description", _get_val(s, "instructions", "")))[:100], _get_val(s, "active_ingredients", _get_val(s, "key_actives", "")) or "—", _get_val(s, "frequency", "Weekly"), _get_val(s, "caution_notes", "") or "—"] for i, s in enumerate(weekly)]
        _write_data_rows(ws3, 2, rows3)
        _auto_fit_columns(ws3)

        if seasonal and isinstance(seasonal, list):
            ws4 = wb.create_sheet("Seasonal Recommendations")
            _write_headers(ws4, 1, ["Season", "Title", "Description", "Tip"])
            sr_rows = [[_get_val(r, "season", ""), _get_val(r, "title", ""), str(_get_val(r, "description", ""))[:100], _get_val(r, "tip", "") or "—"] for r in seasonal]
            _write_data_rows(ws4, 2, sr_rows)
            _auto_fit_columns(ws4)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(
        styles, "Skincare Routine Report",
        f"Patient: {user_name} | Season: {_get_val(routine, 'season', '—')}"
    )

    step_headers = ["#", "Step Title", "Category", "Ingredients", "Frequency"]
    col_w = [0.4*inch, 2*inch, 1.2*inch, 2*inch, 1*inch]

    for label, step_list in [("☀️ Morning Routine", morning), ("🌙 Evening Routine", evening), ("✨ Weekly Treatments", weekly)]:
        elements.append(Paragraph(label, styles["SectionHeading"]))
        if step_list:
            rows = [
                [str(i+1), _get_val(s, "step_title", _get_val(s, "product_name", "")), _get_val(s, "category", _get_val(s, "product_type", "")), _get_val(s, "active_ingredients", _get_val(s, "key_actives", "")) or "—", _get_val(s, "frequency", "Daily")]
                for i, s in enumerate(step_list)
            ]
            elements.append(_build_data_table(step_headers, rows, col_widths=col_w))
        else:
            elements.append(Paragraph("No steps recorded.", styles["Normal"]))
        elements.append(Spacer(1, 8))

    if seasonal and isinstance(seasonal, list):
        elements.append(Paragraph("Seasonal Recommendations", styles["SectionHeading"]))
        sr_headers = ["Season", "Title", "Description"]
        sr_rows = [[_get_val(r, "season", ""), _get_val(r, "title", ""), str(_get_val(r, "description", ""))[:80]] for r in seasonal]
        elements.append(_build_data_table(sr_headers, sr_rows, col_widths=[1.2*inch, 2*inch, 3.4*inch]))
    elif isinstance(seasonal, str) and seasonal.strip():
        elements.append(Paragraph("Seasonal Recommendations", styles["SectionHeading"]))
        elements.append(Paragraph(seasonal, styles["Normal"]))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_product_recommendation_report(user, profile, products, format: str = "pdf") -> bytes:
    """
    Product Recommendation Report: personalized product suggestions.
    """
    user_name = user.name or user.email

    if format == "excel":
        wb, ws = _create_workbook_with_title(
            "Product Recommendation Report",
            f"Patient: {user_name}"
        )
        ws.title = "Recommendations"
        _write_headers(ws, 4, ["#", "Product", "Brand", "Category", "Suitability", "Key Ingredients", "Price (₹)"])
        rows = []
        for i, p in enumerate(products[:50], 1):
            rows.append([
                i, p.get("name", "—"), p.get("brand", "—"),
                p.get("category_label", p.get("category", "—")),
                f"{p.get('suitability_score', 0)}%",
                p.get("key_ingredients_text", "—"),
                p.get("price", "—"),
            ])
        _write_data_rows(ws, 5, rows)
        _auto_fit_columns(ws)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "Product Recommendation Report", f"Patient: {user_name}")

    elements.append(Paragraph("Personalized Product Recommendations", styles["SectionHeading"]))
    if products:
        headers = ["#", "Product", "Brand", "Category", "Suitability"]
        col_w = [0.4*inch, 2.2*inch, 1.2*inch, 1.5*inch, 1*inch]
        rows = [
            [str(i+1), p.get("name", "—"), p.get("brand", "—"),
             p.get("category_label", p.get("category", "—")),
             f"{p.get('suitability_score', 0)}%"]
            for i, p in enumerate(products[:30])
        ]
        elements.append(_build_data_table(headers, rows, col_widths=col_w))
    else:
        elements.append(Paragraph("No product recommendations available. Complete your skin profile to receive personalized suggestions.", styles["Normal"]))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_progress_report(user, assessments, adherence_stats, format: str = "pdf") -> bytes:
    """
    Progress Report: score trends, adherence, streaks, improvement analysis.
    """
    user_name = user.name or user.email

    if format == "excel":
        wb, ws = _create_workbook_with_title("Progress Report", f"Patient: {user_name}")
        ws.title = "Adherence Summary"

        row = 4
        _write_headers(ws, row, ["Metric", "Value"])
        row += 1
        summary_data = [
            ["Total Check-ins", adherence_stats.get("total_checkins", 0)],
            ["Morning Completed", adherence_stats.get("morning_completed", 0)],
            ["Evening Completed", adherence_stats.get("evening_completed", 0)],
            ["Adherence Percentage", f"{adherence_stats.get('adherence_percentage', 0)}%"],
            ["Current Streak", f"{adherence_stats.get('current_streak', 0)} days"],
        ]
        _write_data_rows(ws, row, summary_data)
        _auto_fit_columns(ws)

        ws2 = wb.create_sheet("Score History")
        _write_headers(ws2, 1, ["#", "Date", "Score", "Category", "Risk Level", "Trigger"])
        rows = []
        for i, a in enumerate(assessments[:50], 1):
            rows.append([
                i, _fmt_dt(a.assessment_date), a.skin_health_score,
                a.skin_health_category, a.overall_risk_level, a.trigger_source
            ])
        _write_data_rows(ws2, 2, rows, risk_col=5)
        _auto_fit_columns(ws2)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "Progress Report", f"Patient: {user_name}")

    elements.append(Paragraph("Routine Adherence Summary", styles["SectionHeading"]))
    adh_kv = [
        ("Total Check-ins", str(adherence_stats.get("total_checkins", 0))),
        ("Morning Sessions Completed", str(adherence_stats.get("morning_completed", 0))),
        ("Evening Sessions Completed", str(adherence_stats.get("evening_completed", 0))),
        ("Adherence Rate", f"{adherence_stats.get('adherence_percentage', 0)}%"),
        ("Current Streak", f"{adherence_stats.get('current_streak', 0)} days"),
    ]
    elements.append(_build_kv_table(adh_kv))

    if assessments:
        elements.append(Spacer(1, 12))
        elements.append(Paragraph("Score Trend History", styles["SectionHeading"]))
        trend_headers = ["#", "Date", "Score", "Category", "Risk Level"]
        trend_rows = [
            [str(i+1), _fmt_dt(a.assessment_date), str(a.skin_health_score), a.skin_health_category, a.overall_risk_level]
            for i, a in enumerate(assessments[:20])
        ]
        elements.append(_build_data_table(trend_headers, trend_rows, col_widths=[0.5*inch, 1.8*inch, 1*inch, 1.5*inch, 1.5*inch]))

        if len(assessments) >= 2:
            first = assessments[-1].skin_health_score
            latest = assessments[0].skin_health_score
            diff = latest - first
            direction = "↑ Improved" if diff > 0 else ("↓ Declined" if diff < 0 else "→ Stable")
            elements.append(Spacer(1, 10))
            elements.append(Paragraph(
                f"Overall Trend: {direction} ({'+' if diff > 0 else ''}{diff} points from {first} → {latest})",
                styles["SubSection"]
            ))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_skin_health_report(user, profile, assessments, adherence_stats, routine, format: str = "pdf") -> bytes:
    """
    Comprehensive Skin Health Report: combines profile + score + risks + routine + adherence.
    """
    user_name = _get_val(user, "name") or _get_val(user, "email")
    latest = assessments[0] if assessments else None
    steps = _get_val(routine, "steps", [])
    if not steps:
        m_list = _get_val(routine, "morning_steps", []) or []
        e_list = _get_val(routine, "evening_steps", []) or []
        w_list = _get_val(routine, "weekly_steps", []) or []
        combined = []
        for s in m_list:
            item = dict(s) if isinstance(s, dict) else s
            if isinstance(item, dict):
                item["time_of_day"] = "morning"
            combined.append(item)
        for s in e_list:
            item = dict(s) if isinstance(s, dict) else s
            if isinstance(item, dict):
                item["time_of_day"] = "evening"
            combined.append(item)
        for s in w_list:
            item = dict(s) if isinstance(s, dict) else s
            if isinstance(item, dict):
                item["time_of_day"] = "weekly"
            combined.append(item)
        steps = combined

    if format == "excel":
        wb, ws = _create_workbook_with_title("Comprehensive Skin Health Report", f"Patient: {user_name}")

        # Profile sheet
        ws.title = "Skin Profile"
        row = 4
        profile_data = _get_user_profile_data(user, profile)
        _write_headers(ws, row, ["Field", "Value"])
        _write_data_rows(ws, row + 1, profile_data)
        _auto_fit_columns(ws)

        # Score sheet
        ws2 = wb.create_sheet("Health Score & Risks")
        ws2.merge_cells("A1:B1")
        ws2["A1"] = f"Health Score: {_get_val(latest, 'skin_health_score', '—') if latest else '—'} ({_get_val(latest, 'skin_health_category', '—') if latest else '—'})"
        ws2["A1"].font = EXCEL_TITLE_FONT
        risks = _get_val(latest, "risks", []) if latest else []
        if risks:
            _write_headers(ws2, 3, ["Risk Title", "Risk Level", "Description", "Recommendation"])
            risk_rows = [[_get_val(r, "risk_title", ""), _get_val(r, "risk_level", ""), _get_val(r, "description", ""), _get_val(r, "recommendation", "")] for r in risks]
            _write_data_rows(ws2, 4, risk_rows, risk_col=2)
        _auto_fit_columns(ws2)

        # Routine sheet
        ws3 = wb.create_sheet("Routine Summary")
        _write_headers(ws3, 1, ["Time", "Step", "Category", "Frequency"])
        r_rows = [[_get_val(s, "time_of_day", "").title(), _get_val(s, "step_title", _get_val(s, "product_name", "")), _get_val(s, "category", _get_val(s, "product_type", "")), _get_val(s, "frequency", "Daily")] for s in steps if _get_val(s, "is_active", True)]
        _write_data_rows(ws3, 2, r_rows)
        _auto_fit_columns(ws3)

        # Adherence sheet
        ws4 = wb.create_sheet("Adherence")
        _write_headers(ws4, 1, ["Metric", "Value"])
        adh_rows = [
            ["Adherence Rate", f"{adherence_stats.get('adherence_percentage', 0)}%"],
            ["Current Streak", f"{adherence_stats.get('current_streak', 0)} days"],
            ["Total Check-ins", str(adherence_stats.get("total_checkins", 0))],
        ]
        _write_data_rows(ws4, 2, adh_rows)
        _auto_fit_columns(ws4)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "Comprehensive Skin Health Report", f"Patient: {user_name}")

    # Profile
    elements.append(Paragraph("Skin Profile", styles["SectionHeading"]))
    elements.append(_build_kv_table(_get_user_profile_data(user, profile)))

    # Score & Risks
    if latest:
        elements.append(Spacer(1, 10))
        elements.append(Paragraph("Health Score & Risk Assessment", styles["SectionHeading"]))
        score_kv = [
            ("Health Score", str(_get_val(latest, "skin_health_score", 0))),
            ("Category", str(_get_val(latest, "skin_health_category", "Fair"))),
            ("Overall Risk", str(_get_val(latest, "overall_risk_level", "Low"))),
            ("Assessment Date", _fmt_dt(_get_val(latest, "assessment_date", None))),
        ]
        elements.append(_build_kv_table(score_kv))

        risks = _get_val(latest, "risks", [])
        if risks:
            elements.append(Spacer(1, 6))
            elements.append(Paragraph("Risk Factors", styles["SubSection"]))
            risk_headers = ["Risk", "Level", "Recommendation"]
            risk_rows = [[_get_val(r, "risk_title", ""), _get_val(r, "risk_level", ""), str(_get_val(r, "recommendation", ""))[:80]] for r in risks]
            elements.append(_build_data_table(risk_headers, risk_rows, col_widths=[2*inch, 1*inch, 3.6*inch]))

    # Routine summary
    active_steps = [s for s in steps if _get_val(s, "is_active", True)]
    if active_steps:
        elements.append(Spacer(1, 10))
        elements.append(Paragraph("Active Routine Summary", styles["SectionHeading"]))
        r_headers = ["Time", "Step", "Category", "Frequency"]
        r_rows = [[_get_val(s, "time_of_day", "").title(), _get_val(s, "step_title", _get_val(s, "product_name", "")), _get_val(s, "category", _get_val(s, "product_type", "")), _get_val(s, "frequency", "Daily")] for s in active_steps[:20]]
        elements.append(_build_data_table(r_headers, r_rows, col_widths=[1.2*inch, 2.2*inch, 1.5*inch, 1.5*inch]))

    # Adherence
    elements.append(Spacer(1, 10))
    elements.append(Paragraph("Routine Adherence", styles["SectionHeading"]))
    adh_kv = [
        ("Adherence Rate", f"{adherence_stats.get('adherence_percentage', 0)}%"),
        ("Current Streak", f"{adherence_stats.get('current_streak', 0)} days"),
        ("Total Check-ins", str(adherence_stats.get("total_checkins", 0))),
    ]
    elements.append(_build_kv_table(adh_kv))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_clinical_treatment_report(user, treatments, format: str = "pdf") -> bytes:
    """
    Clinical Treatment Report: all prescriptions and treatment plans for a patient.
    """
    user_name = _get_val(user, "name") or _get_val(user, "email")

    if format == "excel":
        wb, ws = _create_workbook_with_title("Clinical Treatment Report", f"Patient: {user_name}")
        ws.title = "Treatments"
        _write_headers(ws, 4, ["#", "Diagnosis", "Treatment Plan", "Medications/Actives", "Urgency", "Notes", "Date"])
        rows = []
        for i, t in enumerate(treatments, 1):
            rows.append([
                i, _get_val(t, "diagnosis_title", ""), str(_get_val(t, "treatment_plan", ""))[:100],
                _get_val(t, "medications_or_actives", "") or "—", _get_val(t, "urgency_level", "Routine"),
                _get_val(t, "clinical_notes", "") or "—",
                _fmt_dt(_get_val(t, "created_at")),
            ])
        _write_data_rows(ws, 5, rows, risk_col=5)
        _auto_fit_columns(ws)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "Clinical Treatment Report", f"Patient: {user_name}", role_label="Dermatologist")

    elements.append(Paragraph("Prescribed Treatments & Medications", styles["SectionHeading"]))
    if treatments:
        for i, t in enumerate(treatments[:20], 1):
            elements.append(Paragraph(f"{i}. {_get_val(t, 'diagnosis_title', '')}", styles["SubSection"]))
            t_kv = [
                ("Treatment Plan", str(_get_val(t, "treatment_plan", ""))[:200]),
                ("Medications / Actives", _safe_str(_get_val(t, "medications_or_actives", ""))),
                ("Urgency Level", str(_get_val(t, "urgency_level", "Routine"))),
                ("Clinical Notes", _safe_str(_get_val(t, "clinical_notes", ""))),
                ("Date Issued", _fmt_dt(_get_val(t, "created_at"))),
            ]
            elements.append(_build_kv_table(t_kv))
            elements.append(Spacer(1, 6))
    else:
        elements.append(Paragraph("No clinical treatments have been prescribed yet.", styles["Normal"]))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════════
# ADMIN-SPECIFIC REPORTS
# ═══════════════════════════════════════════════════════════════════════════════


def build_platform_summary_report(summary_data: List[dict], format: str = "pdf") -> bytes:
    """Platform Executive Summary for Admin."""
    if format == "excel":
        wb, ws = _create_workbook_with_title("Platform Executive Summary")
        ws.title = "Summary"
        _write_headers(ws, 4, ["Metric", "Value"])
        rows = [[d.get("Metric", ""), d.get("Value", "")] for d in summary_data]
        _write_data_rows(ws, 5, rows)
        _auto_fit_columns(ws)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "Platform Executive Summary", "Administration Report", role_label="Admin")
    elements.append(Paragraph("Platform KPIs", styles["SectionHeading"]))
    kv_data = [(d.get("Metric", ""), str(d.get("Value", ""))) for d in summary_data]
    elements.append(_build_kv_table(kv_data))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_user_directory_report(users_data: List[dict], format: str = "pdf") -> bytes:
    """User & Specialist Directory for Admin."""
    if format == "excel":
        wb, ws = _create_workbook_with_title("User & Specialist Directory")
        ws.title = "Directory"
        _write_headers(ws, 4, ["ID", "Name", "Email", "Status", "Registered"])
        rows = [[d.get("ID", ""), d.get("Name", ""), d.get("Email", ""), d.get("Status", ""), d.get("Registered", "")] for d in users_data]
        _write_data_rows(ws, 5, rows)
        _auto_fit_columns(ws)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "User & Specialist Directory", "Platform Registry", role_label="Admin")
    elements.append(Paragraph("All Registered Accounts", styles["SectionHeading"]))
    if users_data:
        headers = ["ID", "Name", "Email", "Status", "Registered"]
        rows = [[str(d.get("ID", "")), d.get("Name", ""), d.get("Email", ""), d.get("Status", ""), d.get("Registered", "")] for d in users_data]
        elements.append(_build_data_table(headers, rows, col_widths=[0.5*inch, 1.5*inch, 2.2*inch, 1*inch, 1.3*inch]))
    else:
        elements.append(Paragraph("No user records found.", styles["Normal"]))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_assessment_audit_report(assessments_data: List[dict], format: str = "pdf") -> bytes:
    """Clinical Assessment Audit for Admin."""
    if format == "excel":
        wb, ws = _create_workbook_with_title("Clinical Assessment Audit Log")
        ws.title = "Assessments"
        _write_headers(ws, 4, ["ID", "UserID", "Score", "Category", "Risk Level", "Trigger", "Date"])
        rows = [[d.get("AssessmentID", ""), d.get("UserID", ""), d.get("HealthScore", ""), d.get("Category", ""),
                 d.get("RiskLevel", ""), d.get("TriggerSource", ""), d.get("Date", "")] for d in assessments_data]
        _write_data_rows(ws, 5, rows, risk_col=5)
        _auto_fit_columns(ws)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "Clinical Assessment Audit Log", "Full Assessment History", role_label="Admin")
    elements.append(Paragraph("Assessment Records", styles["SectionHeading"]))
    if assessments_data:
        headers = ["ID", "User", "Score", "Category", "Risk", "Date"]
        rows = [
            [str(d.get("AssessmentID", "")), str(d.get("UserID", "")), str(d.get("HealthScore", "")),
             d.get("Category", ""), d.get("RiskLevel", ""), d.get("Date", "")]
            for d in assessments_data[:50]
        ]
        elements.append(_build_data_table(headers, rows, col_widths=[0.7*inch, 0.7*inch, 0.8*inch, 1.2*inch, 1*inch, 1.8*inch]))
    else:
        elements.append(Paragraph("No assessment records found.", styles["Normal"]))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_recommendation_audit_report(recs_data: List[dict], format: str = "pdf") -> bytes:
    """Recommendation History Audit for Admin."""
    if format == "excel":
        wb, ws = _create_workbook_with_title("Recommendation History Audit")
        ws.title = "Recommendations"
        _write_headers(ws, 4, ["ID", "PatientID", "DermatologistID", "Diagnosis", "Treatment", "Actives", "Urgency", "Date"])
        rows = [[d.get("ID", ""), d.get("PatientID", ""), d.get("DermatologistID", ""), d.get("Diagnosis", ""),
                 d.get("TreatmentPlan", "")[:80], d.get("Actives", ""), d.get("Urgency", ""), d.get("Date", "")] for d in recs_data]
        _write_data_rows(ws, 5, rows, risk_col=7)
        _auto_fit_columns(ws)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # PDF
    styles = _get_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=30, bottomMargin=30, leftMargin=40, rightMargin=40)
    elements = _build_header(styles, "Recommendation History Audit", "Clinical Prescriptions & Treatment Trails", role_label="Admin")
    elements.append(Paragraph("Issued Recommendations", styles["SectionHeading"]))
    if recs_data:
        headers = ["ID", "Patient", "Diagnosis", "Urgency", "Date"]
        rows = [
            [str(d.get("ID", "")), str(d.get("PatientID", "")), d.get("Diagnosis", ""), d.get("Urgency", ""), d.get("Date", "")]
            for d in recs_data[:50]
        ]
        elements.append(_build_data_table(headers, rows, col_widths=[0.6*inch, 0.8*inch, 2.5*inch, 1.2*inch, 1.5*inch]))
    else:
        elements.append(Paragraph("No recommendation records found.", styles["Normal"]))

    elements.append(Spacer(1, 20))
    elements.append(_build_footer())
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()
