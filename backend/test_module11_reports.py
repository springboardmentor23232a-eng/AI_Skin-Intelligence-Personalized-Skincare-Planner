import io
from datetime import datetime, timezone, timedelta
import openpyxl
from fastapi import HTTPException

from app.database import SessionLocal, engine, Base
from app.models import (
    User, 
    SkinAssessment, 
    SkinConcern, 
    RiskFactor, 
    Routine, 
    RoutineItem, 
    RoutineProfile, 
    DailyChecklistLog, 
    SkinHealthScoreRecord,
    Product
)
from app.services import report_service
from app.routers.reports import resolve_target_user


def setup_module():
    Base.metadata.create_all(bind=engine)


def make_test_profile(user_id: int, skin_type: str = "Oily", goal: str = "Acne Care") -> RoutineProfile:
    return RoutineProfile(
        user_id=user_id,
        age_group="25-34",
        skin_type=skin_type,
        sensitivity="Normal",
        concerns=["Acne", "Oiliness"],
        acne_severity="Mild",
        oiliness="High",
        dryness="Normal",
        redness_frequency="Rarely",
        has_routine="Yes",
        current_products=["Cleanser"],
        routine_frequency="Daily",
        skincare_irritation="No",
        active_ingredients=["Salicylic Acid"],
        sleep_hours="8",
        water_intake="8",
        stress_level="Low",
        exercise_frequency="Moderate",
        outdoor_hours="1",
        climate="Moderate",
        pollution_exposure="Low",
        sunlight_exposure="Moderate",
        has_allergies="No",
        avoid_ingredients=None,
        has_allergic_reaction="No",
        skincare_time="Quick",
        routine_preference="Minimal",
        budget="Budget",
        skincare_goal=goal
    )


def test_module11_skin_assessment_report_and_pdf():
    db = SessionLocal()
    try:
        # Create test user
        ts = int(datetime.now(timezone.utc).timestamp() * 1000)
        user = User(
            email=f"m11_user_assess_{ts}@example.com",
            name="Assessment Test User",
            role="USER"
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        # 1. Test Empty State
        empty_data = report_service.get_skin_assessment_report_data(user.id, db)
        assert empty_data["has_data"] is False
        assert "No skin assessment" in empty_data["message"]

        # 2. Add real assessment data
        assessment = SkinAssessment(
            user_id=user.id,
            assessment_date=datetime.now(timezone.utc),
            overall_condition="Combination & Mild Acne",
            skin_health_score=78,
            notes="Routine sebum audit conducted."
        )
        db.add(assessment)
        db.commit()
        db.refresh(assessment)

        concern = SkinConcern(
            assessment_id=assessment.id,
            concern_name="Acne & Blemishes",
            severity=0.65,
            priority="High"
        )
        risk = RiskFactor(
            assessment_id=assessment.id,
            risk_name="UV Sensitivity",
            risk_level="Medium",
            description="Elevated photodamage potential."
        )
        db.add_all([concern, risk])
        db.commit()

        # 3. Test Real Data Aggregation
        data = report_service.get_skin_assessment_report_data(user.id, db)
        assert data["has_data"] is True
        assert data["overall_condition"] == "Combination & Mild Acne"
        assert data["skin_health_score"] == 78
        assert len(data["concerns"]) == 1
        assert data["concerns"][0]["name"] == "Acne & Blemishes"
        assert len(data["risks"]) == 1
        assert "medical diagnosis" in data["disclaimer"].lower()

        # 4. Test PDF Generation
        pdf_bytes = report_service.generate_pdf_report(data, "skin_assessment")
        assert isinstance(pdf_bytes, bytes)
        assert pdf_bytes.startswith(b"%PDF-")
        assert len(pdf_bytes) > 1000

        # Clean up
        db.delete(concern)
        db.delete(risk)
        db.delete(assessment)
        db.delete(user)
        db.commit()
    finally:
        db.close()


def test_module11_routine_report_and_excel():
    db = SessionLocal()
    try:
        ts = int(datetime.now(timezone.utc).timestamp() * 1000)
        user = User(
            email=f"m11_user_routine_{ts}@example.com",
            name="Routine Test User",
            role="USER"
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        # 1. Profile & Routine
        profile = make_test_profile(user.id, skin_type="Oily", goal="Pore Clarification")
        db.add(profile)
        db.commit()

        routine = Routine(user_id=user.id, profile_id=profile.id)
        db.add(routine)
        db.commit()
        db.refresh(routine)

        item_am = RoutineItem(
            routine_id=routine.id,
            routine_type="MORNING",
            name="Salicylic Acid Gel Cleanser",
            step_order=1,
            category="Cleanser",
            frequency="Daily AM",
            description="Lather gently with lukewarm water.",
            is_enabled=True
        )
        item_pm = RoutineItem(
            routine_id=routine.id,
            routine_type="EVENING",
            name="Niacinamide 10% Night Serum",
            step_order=1,
            category="Treatment",
            frequency="Daily PM",
            description="Apply 3 drops before moisturizer.",
            is_enabled=True
        )
        db.add_all([item_am, item_pm])
        db.commit()

        # 2. Daily Checklist Logs
        log1 = DailyChecklistLog(
            user_id=user.id,
            logged_at=datetime.now(timezone.utc),
            completed_count=2,
            total_count=2,
            completion_rate=1.0
        )
        db.add(log1)
        db.commit()

        # 3. Test Report Aggregation
        data = report_service.get_routine_report_data(user.id, db)
        assert data["has_data"] is True
        assert len(data["morning_routine"]) == 1
        assert data["morning_routine"][0]["name"] == "Salicylic Acid Gel Cleanser"
        assert len(data["evening_routine"]) == 1

        # 4. Test PDF Generation
        pdf_bytes = report_service.generate_pdf_report(data, "routine")
        assert pdf_bytes.startswith(b"%PDF-")

        # 5. Test Excel Generation
        excel_bytes = report_service.generate_routine_excel(data)
        assert len(excel_bytes) > 500
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        assert "Morning Routine (AM)" in sheet_names
        assert "Evening Routine (PM)" in sheet_names
        assert "Adherence & Profile" in sheet_names

        # Clean up
        db.delete(log1)
        db.delete(item_am)
        db.delete(item_pm)
        db.delete(routine)
        db.delete(profile)
        db.delete(user)
        db.commit()
    finally:
        db.close()


def test_module11_product_recommendations_report_and_excel():
    db = SessionLocal()
    try:
        ts = int(datetime.now(timezone.utc).timestamp() * 1000)
        user = User(
            email=f"m11_user_prod_{ts}@example.com",
            name="Products Test User",
            role="USER"
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        # Profile for recommendation service
        profile = make_test_profile(user.id, skin_type="Dry", goal="Hydration")
        db.add(profile)
        db.commit()

        # Test Aggregation
        data = report_service.get_product_recommendations_report_data(user.id, db)
        assert data["has_data"] is True
        assert len(data["products"]) > 0

        # Test PDF & Excel
        pdf_bytes = report_service.generate_pdf_report(data, "products")
        assert pdf_bytes.startswith(b"%PDF-")

        excel_bytes = report_service.generate_products_excel(data)
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        assert "Recommended Formulations" in wb.sheetnames

        # Clean up
        db.delete(profile)
        db.delete(user)
        db.commit()
    finally:
        db.close()


def test_module11_progress_and_skin_health_reports():
    db = SessionLocal()
    try:
        ts = int(datetime.now(timezone.utc).timestamp() * 1000)
        user = User(
            email=f"m11_user_prog_{ts}@example.com",
            name="Progress Test User",
            role="USER"
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        # Add Score records
        now = datetime.now(timezone.utc)
        rec1 = SkinHealthScoreRecord(
            user_id=user.id,
            calculated_at=now - timedelta(days=14),
            overall_score=70,
            condition_score=70.0,
            lifestyle_score=70.0,
            sleep_score=70.0,
            routine_score=70.0,
            hydration_score=70.0
        )
        rec2 = SkinHealthScoreRecord(
            user_id=user.id,
            calculated_at=now,
            overall_score=82,
            condition_score=85.0,
            lifestyle_score=80.0,
            sleep_score=80.0,
            routine_score=85.0,
            hydration_score=80.0
        )
        db.add_all([rec1, rec2])
        db.commit()

        # Test Progress Report
        prog_data = report_service.get_progress_report_data(user.id, db)
        assert prog_data["has_data"] is True
        assert len(prog_data["score_history"]) >= 2

        prog_pdf = report_service.generate_pdf_report(prog_data, "progress")
        assert prog_pdf.startswith(b"%PDF-")

        prog_excel = report_service.generate_progress_excel(prog_data)
        wb_prog = openpyxl.load_workbook(io.BytesIO(prog_excel))
        assert "Skin Health Score History" in wb_prog.sheetnames

        # Test Skin Health Report
        health_data = report_service.get_skin_health_report_data(user.id, db)
        assert health_data["has_data"] is True
        assert health_data["overall_score"] == 82
        assert len(health_data["components"]) == 5

        health_pdf = report_service.generate_pdf_report(health_data, "skin_health")
        assert health_pdf.startswith(b"%PDF-")

        health_excel = report_service.generate_skin_health_excel(health_data)
        wb_health = openpyxl.load_workbook(io.BytesIO(health_excel))
        assert "Component Breakdown" in wb_health.sheetnames

        # Clean up
        db.delete(rec1)
        db.delete(rec2)
        db.delete(user)
        db.commit()
    finally:
        db.close()


def test_module11_rbac_and_isolation():
    db = SessionLocal()
    try:
        ts = int(datetime.now(timezone.utc).timestamp() * 1000)
        user_a = User(
            email=f"user_a_{ts}@example.com",
            name="User A",
            role="USER"
        )
        user_b = User(
            email=f"user_b_{ts}@example.com",
            name="User B",
            role="USER"
        )
        consultant = User(
            email=f"consultant_{ts}@example.com",
            name="Consultant C",
            role="CONSULTANT"
        )
        db.add_all([user_a, user_b, consultant])
        db.commit()
        db.refresh(user_a)
        db.refresh(user_b)
        db.refresh(consultant)

        # 1. User A accessing own report -> ALLOWED
        target = resolve_target_user(user_a, None, db)
        assert target.id == user_a.id

        target_explicit = resolve_target_user(user_a, user_a.id, db)
        assert target_explicit.id == user_a.id

        # 2. User A accessing User B's report -> FORBIDDEN (403)
        try:
            resolve_target_user(user_a, user_b.id, db)
            assert False, "Expected 403 Forbidden"
        except HTTPException as exc_info:
            assert exc_info.status_code == 403

        # 3. Consultant accessing User A's report -> ALLOWED
        target_client = resolve_target_user(consultant, user_a.id, db)
        assert target_client.id == user_a.id

        # 4. Consultant accessing non-existent user -> NOT FOUND (404)
        try:
            resolve_target_user(consultant, 99999999, db)
            assert False, "Expected 404 Not Found"
        except HTTPException as exc_404:
            assert exc_404.status_code == 404

        # Clean up
        db.delete(user_a)
        db.delete(user_b)
        db.delete(consultant)
        db.commit()
    finally:
        db.close()


if __name__ == "__main__":
    setup_module()
    print("Testing Module 11 Skin Assessment Report & PDF...")
    test_module11_skin_assessment_report_and_pdf()
    print("PASS: Skin Assessment Report & PDF")

    print("Testing Module 11 Routine Report & Excel...")
    test_module11_routine_report_and_excel()
    print("PASS: Routine Report & Excel")

    print("Testing Module 11 Product Recommendations Report & Excel...")
    test_module11_product_recommendations_report_and_excel()
    print("PASS: Product Recommendations Report & Excel")

    print("Testing Module 11 Progress & Skin Health Reports...")
    test_module11_progress_and_skin_health_reports()
    print("PASS: Progress & Skin Health Reports")

    print("Testing Module 11 RBAC & User Isolation...")
    test_module11_rbac_and_isolation()
    print("PASS: RBAC & User Isolation")

    print("\nALL MODULE 11 TESTS PASSED SUCCESSFULLY!")
