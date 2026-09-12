import sys
import os
import io
import csv
import time
import pytest
import pypdf
import openpyxl
from datetime import datetime, timedelta, timezone
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

_backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _backend_dir not in sys.path:
    sys.path.insert(0, _backend_dir)

from app.main import app
from app.db.session import SessionLocal
from app.models import (
    User, SkinProfile, SkinAssessment, SkincareRoutine,
    ProductRecommendation, Product, SkincareLog, SkinProgressPhoto,
    ReminderSetting
)
from app.auth.service import hash_password

client = TestClient(app)


@pytest.fixture(scope="module")
def setup_consistency_users():
    db: Session = SessionLocal()
    ts = int(time.time() * 1000)
    now = datetime.now(timezone.utc)

    # 1. Product catalog seed
    p1 = db.query(Product).filter(Product.name == "Niacinamide 10% + Zinc 1%").first()
    if not p1:
        p1 = Product(
            brand="The Ordinary",
            name="Niacinamide 10% + Zinc 1%",
            category="Serum",
            price=600.0,
            rating=4.8,
            active_ingredients=["Niacinamide 10%", "Zinc PCA 1%"],
            suitable_skin_types=["Oily", "Combination", "All"],
            suitable_concerns=["Acne", "Oiliness", "Pores"],
            description="High-strength vitamin and mineral blemish formula.",
            usage_instructions="Apply a few drops to face morning and evening before heavier creams.",
            purchase_url="https://example.com/the-ordinary-niacinamide"
        )
        db.add(p1)
        db.commit()
        db.refresh(p1)

    p2 = db.query(Product).filter(Product.name == "Barrier Restorative Ceramides Cream").first()
    if not p2:
        p2 = Product(
            brand="CeraVe",
            name="Barrier Restorative Ceramides Cream",
            category="Moisturizer",
            price=1250.0,
            rating=4.9,
            active_ingredients=["Ceramides NP/AP/EOP", "Hyaluronic Acid"],
            suitable_skin_types=["Dry", "Sensitive", "Normal"],
            suitable_concerns=["Dryness", "Barrier Impairment"],
            description="Essential barrier lipids with MVE technology.",
            usage_instructions="Apply liberally to clean dry skin twice daily.",
            purchase_url="https://example.com/cerave-barrier-cream"
        )
        db.add(p2)
        db.commit()
        db.refresh(p2)

    # 2. Comprehensive User A (Oily / Acne focus)
    user_a_email = f"user_a_canonical_{ts}@skincare.com"
    user_a_pass = "CanonicalPassA123!"
    user_a = User(
        full_name=f"Dr. Eleanor Vance {ts}",
        email=user_a_email,
        password=hash_password(user_a_pass),
        role="USER",
        is_active=1,
        is_blocked=0
    )
    db.add(user_a)
    db.commit()
    db.refresh(user_a)

    prof_a = SkinProfile(
        user_id=user_a.id,
        full_name=user_a.full_name,
        age=27,
        gender="Female",
        skin_type="Combination-Oily",
        skin_tone="Fitzpatrick Type IV",
        concerns=["Acne Breakouts", "T-Zone Oiliness"],
        allergies="Benzoyl Peroxide Sensitivity",
        sensitivities="Artificial Fragrances",
        water_intake=2.8,
        stress_level="Moderate",
        lifestyle="Active",
        sleep_quality="7 Hours Restful",
        climate="Humid Tropical",
        uv_exposure="High Daily UV"
    )
    ass_a = SkinAssessment(
        user_id=user_a.id,
        overall_score=78,
        risk_level="Moderate Risk",
        concern_priority="Sebum Balance",
        acne=45,
        hyperpigmentation=30,
        dryness=20,
        oiliness=70,
        redness=35,
        sensitivity=25,
        wrinkles=10,
        fine_lines=15,
        dark_spots=20,
        uneven_tone=25,
        summary="Elevated T-zone sebum with localized papular congestion.",
        created_at=now - timedelta(days=2)
    )
    rout_a_m = SkincareRoutine(
        user_id=user_a.id,
        routine_type="MORNING",
        title="Clarifying Antioxidant Shield Protocol",
        description="Daytime sebum control and barrier shield",
        steps=[
            {
                "step_number": 1,
                "category": "Cleanser",
                "product_name": "Salicylic Foaming Gel",
                "instructions": "Lather onto damp skin for 60 seconds",
                "frequency": "Daily (Morning)",
                "duration": "1 Minute",
                "precautions": "Avoid eye area",
                "expected_benefits": "Unblocks pore debris without moisture strip"
            },
            {
                "step_number": 2,
                "category": "Serum",
                "product_name": "Niacinamide 10% + Zinc 1%",
                "instructions": "Apply 3-4 drops and gently press",
                "frequency": "Daily (Morning)",
                "duration": "1 Minute",
                "precautions": "Do not combine with L-Ascorbic Acid",
                "expected_benefits": "Regulates sebum production and reduces redness"
            }
        ]
    )
    rout_a_e = SkincareRoutine(
        user_id=user_a.id,
        routine_type="EVENING",
        title="Cellular Night Repair Protocol",
        description="Evening deep renewal",
        steps=[
            {
                "step_number": 1,
                "category": "Treatment",
                "product_name": "Azelaic Acid 10% Suspension",
                "instructions": "Apply pea-sized amount after cleansing",
                "frequency": "Daily (Evening)",
                "duration": "1 Minute",
                "precautions": "Use sun protection next morning",
                "expected_benefits": "Targets persistent acne and post-inflammatory erythema"
            }
        ]
    )
    rec_a = ProductRecommendation(
        user_id=user_a.id,
        budget_tier="MEDIUM",
        overall_match_score=94.5,
        recommended_products=[
            {
                "product": {
                    "id": p1.id,
                    "brand": p1.brand,
                    "name": p1.name,
                    "category": p1.category,
                    "price": p1.price,
                    "active_ingredients": p1.active_ingredients,
                    "suitable_skin_types": p1.suitable_skin_types,
                    "suitable_concerns": p1.suitable_concerns,
                    "usage_instructions": p1.usage_instructions,
                    "purchase_url": p1.purchase_url
                },
                "suitability_score": 94.5,
                "match_reasons": ["Ideal match for Combination-Oily skin", "Targets Sebum Balance"],
                "allergy_warnings": []
            }
        ]
    )
    log_a = SkincareLog(
        user_id=user_a.id,
        routine_type="MORNING",
        completed=1,
        logged_date=now - timedelta(days=1),
        notes="Skin significantly calmer and matte throughout day."
    )
    photo_a = SkinProgressPhoto(
        user_id=user_a.id,
        photo_url="https://example.com/photos/eleanor_week2.jpg",
        notes="Week 2 baseline comparison: visible reduction in redness.",
        associated_assessment_id=ass_a.id,
        logged_at=now - timedelta(days=2)
    )
    rem_a = ReminderSetting(
        user_id=user_a.id,
        reminder_type="ROUTINE_MORNING",
        enabled=1,
        time_of_day="07:30",
        recurrence="DAILY"
    )
    db.add_all([prof_a, ass_a, rout_a_m, rout_a_e, rec_a, log_a, photo_a, rem_a])

    # 3. Comprehensive User B (Dry / Sensitive focus)
    user_b_email = f"user_b_canonical_{ts}@skincare.com"
    user_b_pass = "CanonicalPassB123!"
    user_b = User(
        full_name=f"Jonathan Archer {ts}",
        email=user_b_email,
        password=hash_password(user_b_pass),
        role="USER",
        is_active=1,
        is_blocked=0
    )
    db.add(user_b)
    db.commit()
    db.refresh(user_b)

    prof_b = SkinProfile(
        user_id=user_b.id,
        full_name=user_b.full_name,
        age=42,
        gender="Male",
        skin_type="Dry-Sensitive",
        skin_tone="Fitzpatrick Type II",
        concerns=["Micro-Flaking", "Barrier Dehydration"],
        allergies="Sulfates & Detergents",
        sensitivities="Alcohol Denat",
        water_intake=3.0,
        stress_level="Low",
        climate="Continental Cold Dry",
        uv_exposure="Low UV"
    )
    ass_b = SkinAssessment(
        user_id=user_b.id,
        overall_score=84,
        risk_level="Low Risk",
        concern_priority="Lipid Replenishment",
        acne=10,
        hyperpigmentation=15,
        dryness=68,
        oiliness=15,
        redness=20,
        sensitivity=40,
        summary="Impaired stratum corneum moisture seal requiring ceramide support.",
        created_at=now - timedelta(days=1)
    )
    rout_b = SkincareRoutine(
        user_id=user_b.id,
        routine_type="EVENING",
        title="Intensive Ceramide Barrier Occlusion Protocol",
        description="Barrier restoration balm",
        steps=[
            {
                "step_number": 1,
                "category": "Moisturizer",
                "product_name": "Barrier Restorative Ceramides Cream",
                "instructions": "Warm in palms and press onto face",
                "frequency": "Daily (Evening)",
                "expected_benefits": "Rebuilds intercellular lipid bilayer"
            }
        ]
    )
    rec_b = ProductRecommendation(
        user_id=user_b.id,
        budget_tier="MEDIUM",
        overall_match_score=96.0,
        recommended_products=[
            {
                "product": {
                    "id": p2.id,
                    "brand": p2.brand,
                    "name": p2.name,
                    "category": p2.category,
                    "price": p2.price,
                    "active_ingredients": p2.active_ingredients,
                    "suitable_skin_types": p2.suitable_skin_types,
                    "suitable_concerns": p2.suitable_concerns,
                    "usage_instructions": p2.usage_instructions,
                    "purchase_url": p2.purchase_url
                },
                "suitability_score": 96.0,
                "match_reasons": ["Optimal barrier lipid ratio for Dry-Sensitive skin"],
                "allergy_warnings": []
            }
        ]
    )
    db.add_all([prof_b, ass_b, rout_b, rec_b])

    # 4. Completely Empty User (Zero records)
    empty_email = f"user_empty_canonical_{ts}@skincare.com"
    empty_pass = "CanonicalEmptyPass123!"
    user_empty = User(
        full_name=f"Empty User {ts}",
        email=empty_email,
        password=hash_password(empty_pass),
        role="USER",
        is_active=1,
        is_blocked=0
    )
    db.add(user_empty)
    db.commit()
    db.refresh(user_empty)

    # Auth Tokens
    def get_token(em, pw):
        c = TestClient(app)
        res = c.post("/api/auth/login", json={"email": em, "password": pw})
        assert res.status_code == 200
        return res.json()["access_token"]

    return {
        "user_a": {
            "name": user_a.full_name,
            "email": user_a_email,
            "password": user_a_pass,
            "token": get_token(user_a_email, user_a_pass),
            "skin_type": "Combination-Oily",
            "skin_tone": "Fitzpatrick Type IV",
            "concern": "Acne Breakouts",
            "allergies": "Benzoyl Peroxide Sensitivity",
            "product": "Niacinamide 10% + Zinc 1%",
            "brand": "The Ordinary",
            "actives": "Niacinamide 10%",
            "score": "78%"
        },
        "user_b": {
            "name": user_b.full_name,
            "email": user_b_email,
            "password": user_b_pass,
            "token": get_token(user_b_email, user_b_pass),
            "skin_type": "Dry-Sensitive",
            "skin_tone": "Fitzpatrick Type II",
            "concern": "Micro-Flaking",
            "allergies": "Sulfates & Detergents",
            "product": "Barrier Restorative Ceramides Cream",
            "brand": "CeraVe",
            "actives": "Ceramides NP/AP/EOP",
            "score": "84%"
        },
        "empty": {
            "name": user_empty.full_name,
            "email": empty_email,
            "token": get_token(empty_email, empty_pass)
        }
    }


def test_canonical_cross_format_consistency(setup_consistency_users):
    """
    PHASE 16: Automated Cross-Format Consistency Verification.
    Verifies that PDF, XLSX, and CSV are generated from the SAME canonical
    report dataset and contain identical clinical values for the user.
    """
    a = setup_consistency_users["user_a"]
    headers = {"Authorization": f"Bearer {a['token']}"}

    # Fetch all 3 formats
    pdf_res = client.get("/api/reports/export?format=pdf", headers=headers)
    csv_res = client.get("/api/reports/export?format=csv", headers=headers)
    xlsx_res = client.get("/api/reports/export?format=xlsx", headers=headers)

    assert pdf_res.status_code == 200
    assert csv_res.status_code == 200
    assert xlsx_res.status_code == 200

    # 1. Parse PDF
    pdf_reader = pypdf.PdfReader(io.BytesIO(pdf_res.content))
    pdf_text = "".join(p.extract_text() for p in pdf_reader.pages)

    # 2. Parse CSV
    csv_text = csv_res.content.decode("utf-8-sig")
    csv_rows = list(csv.reader(io.StringIO(csv_text)))
    csv_flat = " ".join(" ".join(r) for r in csv_rows)

    # 3. Parse XLSX
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    xlsx_flat_values = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for val in row:
                if val is not None:
                    xlsx_flat_values.append(str(val))
    xlsx_flat = " ".join(xlsx_flat_values)

    # Critical Values that MUST be consistent across PDF, CSV, and XLSX
    critical_values = [
        a["name"],
        a["email"],
        a["skin_type"],
        a["skin_tone"],
        a["concern"],
        a["allergies"],
        a["score"],
        a["product"],
        a["brand"],
        a["actives"]
    ]

    for val in critical_values:
        assert val in pdf_text, f"Value '{val}' missing from PDF export"
        assert val in csv_flat, f"Value '{val}' missing from CSV export"
        assert val in xlsx_flat, f"Value '{val}' missing from XLSX export"

    # Verify Health Score 5-Factor presence across all 3
    health_factors = [
        "Skin Barrier & Condition",
        "Lifestyle & Daily Rhythm",
        "Rest & Nighttime Recovery",
        "Daily Routine Consistency",
        "Daily Hydration Balance"
    ]
    for factor in health_factors:
        assert factor in pdf_text, f"Health factor '{factor}' missing from PDF"
        assert factor in csv_flat, f"Health factor '{factor}' missing from CSV"
        assert factor in xlsx_flat, f"Health factor '{factor}' missing from XLSX"

    # Verify Non-Diagnostic Disclaimer across all 3
    assert "Wellness Disclaimer" in pdf_text
    assert "Wellness Disclaimer" in csv_flat
    assert "Wellness Disclaimer" in xlsx_flat


def test_strict_multi_tenant_isolation(setup_consistency_users):
    """
    PHASE 17: Strict Multi-Tenant Data Isolation Test.
    User A's report must NEVER contain User B's identifiers, biometrics, or recommendations.
    User B's report must NEVER contain User A's identifiers, biometrics, or recommendations.
    """
    a = setup_consistency_users["user_a"]
    b = setup_consistency_users["user_b"]

    a_hdrs = {"Authorization": f"Bearer {a['token']}"}
    b_hdrs = {"Authorization": f"Bearer {b['token']}"}

    # User A Exports
    pdf_a = pypdf.PdfReader(io.BytesIO(client.get("/api/reports/export?format=pdf", headers=a_hdrs).content))
    text_a_pdf = "".join(p.extract_text() for p in pdf_a.pages)
    csv_a = client.get("/api/reports/export?format=csv", headers=a_hdrs).content.decode("utf-8-sig")

    # User B Exports
    pdf_b = pypdf.PdfReader(io.BytesIO(client.get("/api/reports/export?format=pdf", headers=b_hdrs).content))
    text_b_pdf = "".join(p.extract_text() for p in pdf_b.pages)
    csv_b = client.get("/api/reports/export?format=csv", headers=b_hdrs).content.decode("utf-8-sig")

    # User A receives ONLY User A
    assert a["name"] in text_a_pdf and b["name"] not in text_a_pdf
    assert a["email"] in text_a_pdf and b["email"] not in text_a_pdf
    assert a["skin_type"] in text_a_pdf and b["skin_type"] not in text_a_pdf
    assert a["product"] in text_a_pdf and b["product"] not in text_a_pdf

    assert a["name"] in csv_a and b["name"] not in csv_a
    assert a["product"] in csv_a and b["product"] not in csv_a

    # User B receives ONLY User B
    assert b["name"] in text_b_pdf and a["name"] not in text_b_pdf
    assert b["email"] in text_b_pdf and a["email"] not in text_b_pdf
    assert b["skin_type"] in text_b_pdf and a["skin_type"] not in text_b_pdf
    assert b["product"] in text_b_pdf and a["product"] not in text_b_pdf

    assert b["name"] in csv_b and a["name"] not in csv_b
    assert b["product"] in csv_b and a["product"] not in csv_b


def test_empty_user_graceful_handling(setup_consistency_users):
    """
    PHASE 18: Empty and Partial Data Graceful Handling.
    User with 0 assessments, 0 routines, 0 products, 0 logs, 0 photos.
    Ensures zero crashes, valid file binary signatures, and clear status messages.
    """
    emp = setup_consistency_users["empty"]
    hdrs = {"Authorization": f"Bearer {emp['token']}"}

    # 1. PDF
    pdf_res = client.get("/api/reports/export?format=pdf", headers=hdrs)
    assert pdf_res.status_code == 200
    assert pdf_res.content.startswith(b"%PDF-")
    reader = pypdf.PdfReader(io.BytesIO(pdf_res.content))
    pdf_text = "".join(p.extract_text() for p in reader.pages)
    assert emp["name"] in pdf_text
    assert "No clinical skin assessments" in pdf_text
    assert "No routines configured" in pdf_text
    assert "No product recommendations" in pdf_text
    assert "Wellness Disclaimer" in pdf_text

    # 2. XLSX
    xlsx_res = client.get("/api/reports/export?format=xlsx", headers=hdrs)
    assert xlsx_res.status_code == 200
    assert xlsx_res.content.startswith(b"PK")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    assert len(wb.sheetnames) >= 4
    ws_ass = wb["Skin Assessment"]
    assert "No clinical skin assessments" in str(ws_ass["A2"].value)

    # 3. CSV
    csv_res = client.get("/api/reports/export?format=csv", headers=hdrs)
    assert csv_res.status_code == 200
    assert csv_res.content.startswith(b"\xef\xbb\xbf")
    csv_text = csv_res.content.decode("utf-8-sig")
    assert emp["name"] in csv_text
    assert "No assessments recorded yet." in csv_text
    assert "No routine steps recorded yet." in csv_text
    assert "No product recommendations are available yet." in csv_text


def test_zero_credential_leakage(setup_consistency_users):
    """
    PHASE 19: Security & Privacy Audit.
    Confirms that zero passwords, password hashes, JWT tokens, or internal secrets
    leak into any exported report file.
    """
    a = setup_consistency_users["user_a"]
    hdrs = {"Authorization": f"Bearer {a['token']}"}

    pdf_text = "".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(client.get("/api/reports/export?format=pdf", headers=hdrs).content)).pages)
    csv_text = client.get("/api/reports/export?format=csv", headers=hdrs).content.decode("utf-8-sig")
    xlsx_bytes = client.get("/api/reports/export?format=xlsx", headers=hdrs).content
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    xlsx_text = " ".join(str(cell.value) for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.value is not None)

    for export_text in [pdf_text, csv_text, xlsx_text]:
        assert a["password"] not in export_text
        assert a["token"] not in export_text
        assert "password_hash" not in export_text
        assert "$2b$" not in export_text  # bcrypt hash salt/prefix
        assert "SECRET_KEY" not in export_text
