import sys
import os
import io
import csv
import time
import pytest
import pypdf
import openpyxl
from datetime import datetime, timedelta
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

_backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _backend_dir not in sys.path:
    sys.path.insert(0, _backend_dir)

from app.main import app
from app.db.session import SessionLocal
from app.models import User, SkinProfile, SkinAssessment, SkincareRoutine, SkincareLog, SkinProgressPhoto
from app.auth.service import hash_password

client = TestClient(app)


@pytest.fixture(scope="module")
def setup_report_users():
    db: Session = SessionLocal()
    ts = int(time.time() * 1000)
    now = datetime.utcnow()

    # User Alpha (Oily / Acne)
    alpha_email = f"alpha_patient_{ts}@skincare.com"
    alpha_pass = "AlphaPass123!"
    alpha_user = User(
        full_name=f"Alpha Patient {ts}",
        email=alpha_email,
        password=hash_password(alpha_pass),
        role="USER",
        is_active=1,
        is_blocked=0
    )
    db.add(alpha_user)
    db.commit()
    db.refresh(alpha_user)

    alpha_prof = SkinProfile(
        user_id=alpha_user.id,
        full_name=alpha_user.full_name,
        age=24,
        gender="Female",
        skin_type="Oily",
        skin_tone="Medium",
        concerns=["Severe Acne", "Excess Sebum"],
        allergies="Salicylic Acid (Mild Sensitivity)",
        sensitivities="Fragrance",
        water_intake=2.4,
        stress_level="High"
    )
    alpha_ass = SkinAssessment(
        user_id=alpha_user.id,
        overall_score=58,
        risk_level="Moderate Risk",
        concern_priority="Acne Control",
        acne=72,
        dryness=20,
        redness=60,
        oiliness=85,
        sensitivity=40,
        summary="Inflammatory acne breakouts with elevated sebum secretion.",
        created_at=now - timedelta(days=5)
    )
    alpha_rout = SkincareRoutine(
        user_id=alpha_user.id,
        routine_type="MORNING",
        title="Acne Clarifying Morning Protocol",
        description="Clarifying protocol",
        steps=[
            {"step_number": 1, "category": "Cleanser", "product_name": "Zinc Clarifying Foam", "instructions": "Cleanse gently"},
            {"step_number": 2, "category": "Treatment", "product_name": "Niacinamide 10% Serum", "instructions": "Apply 3 drops"}
        ]
    )
    alpha_log = SkincareLog(
        user_id=alpha_user.id,
        routine_type="MORNING",
        completed=1,
        logged_date=now - timedelta(days=1),
        notes="Skin feeling less oily"
    )
    alpha_photo = SkinProgressPhoto(
        user_id=alpha_user.id,
        photo_url="https://example.com/alpha_day1.jpg",
        notes="Baseline breakout notes",
        logged_at=now - timedelta(days=5)
    )
    db.add_all([alpha_prof, alpha_ass, alpha_rout, alpha_log, alpha_photo])

    # User Beta (Dry / Sensitive)
    beta_email = f"beta_patient_{ts}@skincare.com"
    beta_pass = "BetaPass123!"
    beta_user = User(
        full_name=f"Beta Patient {ts}",
        email=beta_email,
        password=hash_password(beta_pass),
        role="USER",
        is_active=1,
        is_blocked=0
    )
    db.add(beta_user)
    db.commit()
    db.refresh(beta_user)

    beta_prof = SkinProfile(
        user_id=beta_user.id,
        full_name=beta_user.full_name,
        age=38,
        gender="Male",
        skin_type="Dry",
        skin_tone="Fair",
        concerns=["Extreme Dryness", "Barrier Impairment"],
        allergies="Benzoyl Peroxide",
        sensitivities="Alcohol-based toners",
        water_intake=3.2,
        stress_level="Low"
    )
    beta_ass = SkinAssessment(
        user_id=beta_user.id,
        overall_score=86,
        risk_level="Low Risk",
        concern_priority="Moisture Retention",
        acne=10,
        dryness=78,
        redness=15,
        oiliness=12,
        sensitivity=35,
        summary="Dry moisture-depleted barrier requiring intensive lipid replenishment.",
        created_at=now - timedelta(days=3)
    )
    beta_rout = SkincareRoutine(
        user_id=beta_user.id,
        routine_type="EVENING",
        title="Ceramide Barrier Night Balm Protocol",
        description="Barrier restorative night care",
        steps=[
            {"step_number": 1, "category": "Moisturizer", "product_name": "Triple Ceramide Lipid Cream", "instructions": "Apply liberally"}
        ]
    )
    beta_log = SkincareLog(
        user_id=beta_user.id,
        routine_type="EVENING",
        completed=1,
        logged_date=now - timedelta(days=1),
        notes="Flakiness subsided"
    )
    db.add_all([beta_prof, beta_ass, beta_rout, beta_log])

    # User Empty (0 profile, 0 assessments, 0 routines, 0 logs)
    empty_email = f"empty_patient_{ts}@skincare.com"
    empty_pass = "EmptyPass123!"
    empty_user = User(
        full_name=f"Empty Patient {ts}",
        email=empty_email,
        password=hash_password(empty_pass),
        role="USER",
        is_active=1,
        is_blocked=0
    )
    db.add(empty_user)
    db.commit()
    db.refresh(empty_user)

    # Tokens
    def get_token(em, pw):
        c = TestClient(app)
        res = c.post("/api/auth/login", json={"email": em, "password": pw})
        assert res.status_code == 200
        return res.json()["access_token"]

    return {
        "alpha": {
            "id": alpha_user.id,
            "name": alpha_user.full_name,
            "email": alpha_email,
            "password": alpha_pass,
            "token": get_token(alpha_email, alpha_pass),
            "skin_type": "Oily",
            "concern": "Severe Acne",
            "score": 58,
            "product": "Zinc Clarifying Foam"
        },
        "beta": {
            "id": beta_user.id,
            "name": beta_user.full_name,
            "email": beta_email,
            "password": beta_pass,
            "token": get_token(beta_email, beta_pass),
            "skin_type": "Dry",
            "concern": "Extreme Dryness",
            "score": 86,
            "product": "Triple Ceramide Lipid Cream"
        },
        "empty": {
            "id": empty_user.id,
            "name": empty_user.full_name,
            "email": empty_email,
            "password": empty_pass,
            "token": get_token(empty_email, empty_pass)
        }
    }


def test_pdf_export_parseability_and_user_data(setup_report_users):
    alpha = setup_report_users["alpha"]
    headers = {"Authorization": f"Bearer {alpha['token']}"}

    res = client.get("/api/reports/export?format=pdf", headers=headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert 'filename="skin-health-report.pdf"' in res.headers["content-disposition"]
    assert res.content.startswith(b"%PDF"), "Response is not a valid binary PDF stream"

    # Open and parse with pypdf
    reader = pypdf.PdfReader(io.BytesIO(res.content))
    assert len(reader.pages) >= 1

    extracted_text = ""
    for page in reader.pages:
        extracted_text += page.extract_text() + "\n"

    # Verify real authenticated user health data is present
    assert alpha["name"] in extracted_text
    assert alpha["email"] in extracted_text
    assert "Oily" in extracted_text
    assert "Severe Acne" in extracted_text
    assert "58%" in extracted_text
    assert "Acne Control" in extracted_text
    assert "Zinc Clarifying Foam" in extracted_text
    assert "Wellness Disclaimer" in extracted_text

    # Verify no credential leakage
    assert alpha["password"] not in extracted_text
    assert alpha["token"] not in extracted_text


def test_csv_export_parseability_and_user_data(setup_report_users):
    alpha = setup_report_users["alpha"]
    headers = {"Authorization": f"Bearer {alpha['token']}"}

    res = client.get("/api/reports/export?format=csv", headers=headers)
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    assert 'filename="skin-health-report.csv"' in res.headers["content-disposition"]

    # Verify UTF-8 BOM for Excel compatibility
    assert res.content.startswith(b"\xef\xbb\xbf"), "CSV stream missing UTF-8 BOM for Excel"

    # Parse with standard python csv reader
    csv_text = res.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(csv_text)))
    assert len(rows) > 10

    flat_csv = " ".join(" ".join(r) for r in rows)
    assert alpha["name"] in flat_csv
    assert alpha["email"] in flat_csv
    assert "Oily" in flat_csv
    assert "Severe Acne" in flat_csv
    assert "58%" in flat_csv
    assert "Zinc Clarifying Foam" in flat_csv
    assert "Wellness Disclaimer" in flat_csv
    assert alpha["password"] not in flat_csv


def test_xlsx_export_parseability_with_openpyxl(setup_report_users):
    alpha = setup_report_users["alpha"]
    headers = {"Authorization": f"Bearer {alpha['token']}"}

    res = client.get("/api/reports/export?format=xlsx", headers=headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="skin-health-report.xlsx"' in res.headers["content-disposition"]
    assert res.content.startswith(b"PK"), "XLSX must be a genuine ZIP/openpyxl binary archive"

    # Open and parse with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    expected_sheets = ["Health Summary", "Skin Assessment", "Routine Adherence", "Progress History"]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Missing expected worksheet: {s}"

    ws_summary = wb["Health Summary"]
    # Check patient name in Health Summary
    name_found = any(cell.value == alpha["name"] for row in ws_summary.iter_rows() for cell in row)
    assert name_found, "User's actual full name not found in XLSX Health Summary worksheet"

    # Check assessment in Skin Assessment
    ws_ass = wb["Skin Assessment"]
    score_found = any(cell.value == "58%" for row in ws_ass.iter_rows() for cell in row)
    assert score_found, "User's actual score not found in XLSX Skin Assessment worksheet"

    # Verify no passwords or tokens
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for val in row:
                if val:
                    assert alpha["password"] not in str(val)
                    assert alpha["token"] not in str(val)


def test_user_data_isolation_across_reports(setup_report_users):
    alpha = setup_report_users["alpha"]
    beta = setup_report_users["beta"]

    alpha_headers = {"Authorization": f"Bearer {alpha['token']}"}
    beta_headers = {"Authorization": f"Bearer {beta['token']}"}

    # Fetch Alpha's reports
    alpha_pdf = client.get("/api/reports/export?format=pdf", headers=alpha_headers)
    alpha_csv = client.get("/api/reports/export?format=csv", headers=alpha_headers).content.decode("utf-8-sig")

    # Fetch Beta's reports
    beta_pdf = client.get("/api/reports/export?format=pdf", headers=beta_headers)
    beta_csv = client.get("/api/reports/export?format=csv", headers=beta_headers).content.decode("utf-8-sig")

    # Alpha's PDF has Alpha's data, NEVER Beta's data
    alpha_reader = pypdf.PdfReader(io.BytesIO(alpha_pdf.content))
    alpha_text = "".join(p.extract_text() for p in alpha_reader.pages)
    assert alpha["name"] in alpha_text
    assert beta["name"] not in alpha_text
    assert "Zinc Clarifying Foam" in alpha_text
    assert "Triple Ceramide Lipid Cream" not in alpha_text

    # Beta's PDF has Beta's data, NEVER Alpha's data
    beta_reader = pypdf.PdfReader(io.BytesIO(beta_pdf.content))
    beta_text = "".join(p.extract_text() for p in beta_reader.pages)
    assert beta["name"] in beta_text
    assert alpha["name"] not in beta_text
    assert "Triple Ceramide Lipid Cream" in beta_text
    assert "Zinc Clarifying Foam" not in beta_text

    # CSV isolation
    assert alpha["name"] in alpha_csv and beta["name"] not in alpha_csv
    assert beta["name"] in beta_csv and alpha["name"] not in beta_csv


def test_empty_data_handling_in_all_formats(setup_report_users):
    empty = setup_report_users["empty"]
    headers = {"Authorization": f"Bearer {empty['token']}"}

    # 1. PDF with 0 data
    res_pdf = client.get("/api/reports/export?format=pdf", headers=headers)
    assert res_pdf.status_code == 200
    reader = pypdf.PdfReader(io.BytesIO(res_pdf.content))
    pdf_text = "".join(p.extract_text() for p in reader.pages)
    assert empty["name"] in pdf_text
    assert "No clinical skin assessments" in pdf_text
    assert "No routines configured" in pdf_text
    assert "Wellness Disclaimer" in pdf_text

    # 2. CSV with 0 data
    res_csv = client.get("/api/reports/export?format=csv", headers=headers)
    assert res_csv.status_code == 200
    csv_text = res_csv.content.decode("utf-8-sig")
    assert empty["name"] in csv_text
    assert "No assessments recorded yet." in csv_text
    assert "No routine steps recorded yet." in csv_text

    # 3. XLSX with 0 data
    res_xlsx = client.get("/api/reports/export?format=xlsx", headers=headers)
    assert res_xlsx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    assert "Health Summary" in wb.sheetnames
    ws_ass = wb["Skin Assessment"]
    val_msg = ws_ass["A2"].value
    assert "No clinical skin assessments" in val_msg


def test_unauthenticated_and_invalid_formats():
    unauth = TestClient(app)
    # Unauthenticated rejected with 401
    res = unauth.get("/api/reports/export?format=pdf")
    assert res.status_code == 401

    # Invalid format rejected with 422 (or 400 pattern mismatch)
    # Register temporary user
    db = SessionLocal()
    ts = int(time.time() * 1000)
    u = User(full_name="Tester", email=f"tester_{ts}@test.com", password=hash_password("pass"), role="USER", is_active=1, is_blocked=0)
    db.add(u)
    db.commit()
    auth_res = client.post("/api/auth/login", json={"email": u.email, "password": "pass"})
    token = auth_res.json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    res_inv = client.get("/api/reports/export?format=docx", headers=h)
    assert res_inv.status_code in [400, 422]
